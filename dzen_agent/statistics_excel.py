"""Read original Excel bytes without saving, recalculating or changing the workbook."""
from __future__ import annotations

import hashlib
import io
import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from itertools import zip_longest
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from dzen_agent.statistics import CSVPreview, detect_mapping
from dzen_agent.store import dumps

MAX_ROWS = 50000
MAX_COLUMNS = 512


def _open(raw: bytes, *, data_only=True):
    if not raw or len(raw) > 20 * 1024 * 1024:
        raise ValueError("Загрузите непустой XLSX размером до 20 МБ.")
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Excel-файл слишком большой после распаковки (максимум 100 МБ).")
    except BadZipFile as error:
        raise ValueError("Не удалось прочитать XLSX. Проверьте, что это исправная выгрузка Excel без пароля.") from error
    try:
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=data_only, keep_links=False)
    except (BadZipFile, InvalidFileException, KeyError, ParseError, OSError, ValueError) as error:
        raise ValueError("Не удалось прочитать XLSX. Проверьте, что это исправная выгрузка Excel без пароля.") from error


def excel_sheets(raw: bytes) -> list[str]:
    workbook = _open(raw)
    try:
        names = [sheet.title for sheet in workbook.worksheets]
        if not names:
            raise ValueError("В Excel-файле нет листов с таблицами.")
        return names
    finally:
        workbook.close()


def _value(cell, formula_cell):
    value = cell.value
    if formula_cell.data_type == "f" and value is None:
        raise ValueError(f"Ячейка {formula_cell.coordinate}: у формулы нет сохранённого результата. "
                         "Пересчитайте файл в Excel и сохраните его либо используйте выгрузку с готовыми значениями.")
    if cell.data_type == "e":
        raise ValueError(f"Ячейка {cell.coordinate}: ошибка Excel {value}. Исправьте данные перед импортом.")
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel stores 3.5% as .035; CSV exports contain 3.5%. Preserve the same unit.
        fmt = cell.number_format.split(";")[0]
        if "%" in re.sub(r'"[^"]*"|\\.', "", fmt):
            return f"{Decimal(str(value)) * 100}%"
    return value.strip() if isinstance(value, str) else value


def preview_excel(raw: bytes, *, sheet_name: str | None = None, header_row: int | None = None) -> CSVPreview:
    if header_row is not None and (type(header_row) is not int or not 1 <= header_row <= 50):
        raise ValueError("Строка заголовков должна быть от 1 до 50.")
    workbook = _open(raw)
    formula_book = None
    try:
        sheets = [s.title for s in workbook.worksheets]
        if not sheets:
            raise ValueError("В Excel-файле нет листов с таблицами.")
        if sheet_name is None:
            if len(sheets) != 1:
                raise ValueError("Выберите лист Excel для импорта.")
            sheet_name = sheets[0]
        if sheet_name not in sheets:
            raise ValueError("Выбранный лист Excel не найден.")
        formula_book = _open(raw, data_only=False)
        sheet, formula_sheet = workbook[sheet_name], formula_book[sheet_name]
        # Some exporters declare A1:A1 or an entire formatted sheet as the used range.
        sheet.reset_dimensions()
        formula_sheet.reset_dimensions()
        rows, row_numbers = [], []
        for number, (cells, formulas) in enumerate(zip_longest(sheet.iter_rows(), formula_sheet.iter_rows()), 1):
            if number > MAX_ROWS + 100:
                raise ValueError("Слишком много строк Excel: максимум 50 000 строк данных.")
            if cells is None or formulas is None:
                raise ValueError("Не удалось согласованно прочитать значения Excel.")
            values = [_value(c, f) for c, f in zip(cells, formulas, strict=True)]
            while values and values[-1] is None:
                values.pop()
            if len(values) > MAX_COLUMNS:
                raise ValueError("Слишком много колонок Excel: максимум 512.")
            rows.append(values)
            row_numbers.append(number)
        if not rows:
            raise ValueError("Выбранный лист Excel пуст.")
        if header_row is None:
            candidates = []
            for index, row in enumerate(rows[:50]):
                header = [str(value).strip() if value is not None else "" for value in row]
                recognized = len(detect_mapping(header))
                if len(header) > 1 and any(isinstance(v, str) and v.strip() for v in row):
                    candidates.append((recognized * 100 + min(len(header), 50), -index))
            if not candidates:
                raise ValueError("Не найден заголовок Excel. Нужны как минимум две колонки.")
            _, negative_index = max(candidates)
            header_row = -negative_index + 1
        if header_row > len(rows):
            raise ValueError("Указанная строка заголовков отсутствует на листе.")
        data = [(number, row) for number, row in zip(row_numbers[header_row:], rows[header_row:], strict=True)
                if any(value is not None and value != "" for value in row)]
        if not data:
            raise ValueError("На выбранном листе нет строк со статистикой.")
        if len(data) > MAX_ROWS:
            raise ValueError("Слишком много строк: максимум 50 000.")
        width = max(len(rows[header_row - 1]), max(len(row) for _, row in data))
        header = rows[header_row - 1]
        columns = [str(header[i]).strip() if i < len(header) and header[i] not in (None, "")
                   else f"Колонка {i + 1}" for i in range(width)]
        if len(set(columns)) != len(columns):
            raise ValueError("В Excel есть одинаковые названия колонок. Выберите правильную строку заголовков.")
        records = [dict(zip(columns, row + [None] * (width - len(row)), strict=True)) for _, row in data]
        fingerprint = hashlib.sha256(raw + b"\0xlsx:" + dumps([sheet_name, header_row]).encode()).hexdigest()
        return CSVPreview(columns, records, detect_mapping(columns), "xlsx", "", fingerprint,
            source_format="xlsx", sheet_name=sheet_name, header_row=header_row,
            row_numbers=[number for number, _ in data], source_sha256=hashlib.sha256(raw).hexdigest())
    except (BadZipFile, ParseError, KeyError, OSError) as error:
        raise ValueError("Excel-файл повреждён или его лист не удалось прочитать.") from error
    finally:
        workbook.close()
        if formula_book:
            formula_book.close()
